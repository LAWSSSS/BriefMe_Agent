"""永锋废钢：人工检判表与视觉质检实绩对齐。

人工表是底账，一行一趟车。视觉侧要同时满足：质检已完成，并且扣重扣杂结果里
留有明细。扣重为 0、说明是「0.00吨（）」、扣杂和厚度都没有明细的车次
（例如 2026-09-06 冀AZC799）不算检判留存，对齐状态写「未检判」。
同一天、同一车牌对得上的，写入总计和明细。

同一天同一车牌有多趟时，按净重最接近的一对一配对，避免把两趟扣重写到同一行。

CLI：
  python -m agent.yongfeng.scrap_align --manual D:\\BriefMe_Agent\\废钢人工判定.xlsx
"""
from __future__ import annotations

import argparse
import logging
import re
from collections import defaultdict
from dataclasses import dataclass, field
from datetime import date, datetime
from pathlib import Path
from typing import Callable, Iterator

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.utils.datetime import from_excel

from agent.yongfeng.scrap_client import YongfengScrapClient

logger = logging.getLogger(__name__)

STATUS_ALIGNED = "已对齐"
STATUS_MISSING = "未检判"

REASON_MATCHED = "matched"
REASON_NO_VISUAL_DAY = "no_visual_day"
REASON_NO_PLATE = "no_plate"
REASON_INCOMPLETE = "incomplete"
REASON_EXTRA_MANUAL = "extra_manual"
REASON_EMPTY_PLATE = "empty_plate"
REASON_NO_DETAIL = "no_detail"

_EMPTY_DEDUCT_TEXT = re.compile(r"^0+(?:\.0+)?吨（\s*）$")

_STATUS_LABELS = {
    1: "毛重计量完成",
    2: "质检中",
    3: "暂停",
    4: "质检完成",
    5: "皮重计量完成",
}

_HEADER_FIELDS = {
    "日期": "date",
    "货物名称": "cargo",
    "供货厂家": "supplier",
    "车号": "plate",
    "车牌": "plate",
    "车牌号": "plate",
    "重量（吨）": "weight_ton",
    "重量(吨)": "weight_ton",
    "使用单位": "user_unit",
    "尺寸超标": "size_issue",
    "尺寸超标扣吨": "size_ton",
    "厚度不足": "thick_issue",
    "厚度不足扣吨": "thick_ton",
    "非钢杂质": "impurity_issue",
    "非钢杂质扣吨": "impurity_ton",
    "水分|油污": "moisture_issue",
    "水分|油污扣吨": "moisture_ton",
    "水分油污": "moisture_issue",
    "水分油污扣吨": "moisture_ton",
    "降级处理吨位": "downgrade_issue",
    "降级处理": "downgrade_issue",
    "降级处理扣吨": "downgrade_ton",
    "其他问题": "other_issue",
    "其他问题扣吨": "other_ton",
    "处理结果": "result",
    "扣除重量（吨）": "deduct_ton",
    "扣除重量(吨)": "deduct_ton",
}

_XLSX_IN_TEXT = re.compile(
    r"[A-Za-z]:\\(?:[^\\/:*?\"<>|\r\n]+\\)*[^\\/:*?\"<>|\r\n]+\.xlsx"
    r"|/(?:[^/\s\"']+/)*[^/\s\"']+\.xlsx",
    re.IGNORECASE,
)

# |视觉总计 − 人工扣除| ≤ 150kg 记为合格。未检判没有视觉扣重，不进准确率。
DEDUCT_TOLERANCE_KG = 150

_GROUP_FILLS = {
    "人工检判": "E8F0FE",
    "扣重差异": "FFE082",
    "对齐": "FFF3E0",
    "视觉参考": "E8F5E9",
    "扣重扣杂总计": "F3E5F5",
    "扣杂明细": "FFEBEE",
    "厚度扣重明细": "E0F7FA",
    "对比": "FFFDE7",
}

_WRAP_KEYS = {"reason", "supplier", "total_desc", "standard", "thick_vis_desc", "impurity_vis_desc"}
_DIFF_KEYS = {"deduct_ton", "total_ton", "deduct_diff", "deduct_diff_kg", "verdict"}

_COLUMNS: list[tuple[str, str, str, int, str]] = [
    ("人工检判", "人工表行号", "row_no", 12, "int"),
    ("人工检判", "日期", "date", 12, "text"),
    ("人工检判", "车号", "plate", 14, "text"),
    ("人工检判", "货物名称", "cargo", 12, "text"),
    ("人工检判", "供货厂家", "supplier", 28, "text"),
    ("人工检判", "人工重量（吨）", "weight_ton", 14, "ton"),
    ("人工检判", "使用单位", "user_unit", 12, "text"),
    ("人工检判", "尺寸超标", "size_issue", 12, "text"),
    ("人工检判", "尺寸超标扣吨", "size_ton", 14, "ton"),
    ("人工检判", "厚度不足", "thick_issue", 12, "text"),
    ("人工检判", "厚度不足扣吨", "thick_ton", 14, "ton"),
    ("人工检判", "非钢杂质", "impurity_issue", 12, "text"),
    ("人工检判", "非钢杂质扣吨", "impurity_ton", 14, "ton"),
    ("人工检判", "水分油污", "moisture_issue", 12, "text"),
    ("人工检判", "水分油污扣吨", "moisture_ton", 14, "ton"),
    ("人工检判", "降级处理", "downgrade_issue", 22, "text"),
    ("人工检判", "降级处理扣吨", "downgrade_ton", 14, "ton"),
    ("人工检判", "其他问题", "other_issue", 16, "text"),
    ("人工检判", "其他问题扣吨", "other_ton", 14, "ton"),
    ("人工检判", "处理结果", "result", 16, "text"),
    ("扣重差异", "人工扣除重量（吨）", "deduct_ton", 18, "ton"),
    ("扣重差异", "视觉总计重量（吨）", "total_ton", 18, "ton"),
    ("扣重差异", "扣重差值（吨）", "deduct_diff", 14, "ton"),
    ("扣重差异", "扣重差值（kg）", "deduct_diff_kg", 14, "kg"),
    ("扣重差异", "扣重判定", "verdict", 12, "text"),
    ("对齐", "对齐状态", "status", 12, "text"),
    ("对齐", "配对说明", "reason", 42, "text"),
    ("视觉参考", "净重（吨）", "net_ton", 12, "ton"),
    ("视觉参考", "毛重（吨）", "gross_ton", 12, "ton"),
    ("视觉参考", "物料名称", "material_name", 22, "text"),
    ("视觉参考", "申报料型", "declared_name", 12, "text"),
    ("扣重扣杂总计", "总计占比（%）", "total_rate", 14, "rate"),
    ("扣重扣杂总计", "总计说明", "total_desc", 36, "text"),
    ("扣重扣杂总计", "计算标准", "standard", 36, "text"),
    ("扣杂明细", "扣杂重量（吨）", "impurity_vis_ton", 14, "ton"),
    ("扣杂明细", "扣杂占比（%）", "impurity_vis_rate", 14, "rate"),
    ("扣杂明细", "扣杂说明", "impurity_vis_desc", 28, "text"),
    ("厚度扣重明细", "厚度扣重重量（吨）", "thick_vis_ton", 16, "ton"),
    ("厚度扣重明细", "厚度扣重占比（%）", "thick_vis_rate", 16, "rate"),
    ("厚度扣重明细", "厚度扣重说明", "thick_vis_desc", 48, "text"),
    ("对比", "其他扣重明细", "other_vis", 28, "text"),
    ("对比", "扣杂差值（吨）", "impurity_diff", 14, "ton"),
    ("对比", "厚度扣重差值（吨）", "thick_diff", 16, "ton"),
]


@dataclass
class ManualRow:
    """人工检判表里的一行。"""

    row_no: int
    date: str
    plate: str
    plate_key: str
    cargo: str = ""
    supplier: str = ""
    weight_ton: float | None = None
    user_unit: str = ""
    size_issue: str = ""
    size_ton: float | None = None
    thick_issue: str = ""
    thick_ton: float | None = None
    impurity_issue: str = ""
    impurity_ton: float | None = None
    moisture_issue: str = ""
    moisture_ton: float | None = None
    downgrade_issue: str = ""
    downgrade_ton: float | None = None
    other_issue: str = ""
    other_ton: float | None = None
    result: str = ""
    deduct_ton: float | None = None


@dataclass
class VisualTrip:
    """视觉质检实绩里的一趟车。checked 为真才算有质检完成记录。"""

    date: str
    plate: str
    plate_key: str
    flow_code: str
    checked: bool
    net_ton: float | None
    gross_ton: float | None
    create_time: str
    raw: dict = field(default_factory=dict)
    # 扣重为 0 且没有扣杂/厚度明细时为 False，不参与对齐。
    retained: bool = True
    deduction: dict = field(default_factory=dict)


@dataclass
class MatchResult:
    manual: ManualRow
    status: str
    reason_code: str
    reason: str
    trip: VisualTrip | None = None
    deduction: dict = field(default_factory=dict)


def normalize_plate(value: object) -> str:
    """车牌比较键：去空白、拉丁字母大写、全角数字字母转半角。"""
    text = str(value or "").strip().upper()
    text = text.replace(" ", "").replace("\u3000", "").replace("-", "")
    wide = "０１２３４５６７８９ＡＢＣＤＥＦＧＨＩＪＫＬＭＮＯＰＱＲＳＴＵＶＷＸＹＺ"
    narrow = "0123456789ABCDEFGHIJKLMNOPQRSTUVWXYZ"
    return text.translate(str.maketrans(wide, narrow))


def is_yongfeng_manual_align(message: str) -> bool:
    """永锋废钢「视觉质检实绩 vs 人工检判表」对齐指令，不抢原图下载和烧结矿。"""
    text = message or ""
    if "盛隆" in text or "镔鑫" in text:
        return False
    if "永锋" not in text:
        return False
    if any(k in text for k in ("检判原图", "智能判级照片", "多标签", "烧结", "颗粒度", "打包带")):
        return False
    mentions_manual = any(k in text for k in ("人工检判", "人工判定", "质检实绩", "扣重扣杂"))
    mentions_align = any(k in text for k in ("对齐", "对照", "筛选"))
    return mentions_manual and mentions_align


def extract_xlsx_path(message: str) -> str:
    """从用户消息里抽出第一个 xlsx 绝对路径。"""
    match = _XLSX_IN_TEXT.search(message or "")
    return match.group(0) if match else ""


def uploaded_to_path(uploaded: object) -> str:
    """把 Gradio 文件控件的返回值收成路径。"""
    if uploaded is None:
        return ""
    if isinstance(uploaded, (str, Path)):
        return str(uploaded).strip()
    if isinstance(uploaded, dict):
        return str(uploaded.get("path") or uploaded.get("name") or "").strip()
    name = getattr(uploaded, "name", "")
    return str(name or "").strip()


def resolve_manual_path(message: str, uploaded: object = None) -> str:
    """消息里的 xlsx 路径优先，否则用界面上选取的文件。"""
    typed = extract_xlsx_path(message)
    if typed:
        return typed
    return uploaded_to_path(uploaded)


def _num(value: object) -> float | None:
    if value is None or value == "":
        return None
    try:
        return float(value)
    except (TypeError, ValueError):
        return None


def _text(value: object) -> str:
    if value is None:
        return ""
    return str(value).strip()


def _sub(left: float | None, right: float | None) -> float | None:
    if left is None or right is None:
        return None
    return round(left - right, 4)


def judge_deduct_diff(diff_ton: float | None) -> str:
    """视觉总计减去人工扣除。绝对值不超过 150kg 为合格，未检判返回空。"""
    if diff_ton is None:
        return ""
    diff_kg = round(abs(diff_ton) * 1000, 1)
    return "合格" if diff_kg <= DEDUCT_TOLERANCE_KG else "不合格"


def kg_to_ton(value: object) -> float | None:
    """过磅重量在接口里是千克。扣重字段本身已是吨，不要走这个函数。"""
    num = _num(value)
    if num is None:
        return None
    if abs(num) >= 1000:
        return num / 1000.0
    return num


def parse_manual_date(value: object) -> str:
    """人工表日期：datetime、Excel 序列号、YYYY-MM-DD 都收成同一天。"""
    if isinstance(value, datetime):
        return value.strftime("%Y-%m-%d")
    if isinstance(value, date):
        return value.strftime("%Y-%m-%d")
    if isinstance(value, (int, float)) and not isinstance(value, bool) and value > 20000:
        return from_excel(value).strftime("%Y-%m-%d")
    text = _text(value)
    if not text:
        raise ValueError("日期为空")
    for fmt in ("%Y-%m-%d", "%Y/%m/%d", "%Y.%m.%d", "%Y-%m-%d %H:%M:%S", "%Y/%m/%d %H:%M:%S"):
        try:
            return datetime.strptime(text[:19] if " " in text else text, fmt).strftime("%Y-%m-%d")
        except ValueError:
            continue
    matched = re.match(r"(\d{4})\D(\d{1,2})\D(\d{1,2})", text)
    if matched:
        return f"{int(matched.group(1)):04d}-{int(matched.group(2)):02d}-{int(matched.group(3)):02d}"
    raise ValueError(f"无法识别日期: {value!r}")


def _status_label(status: object) -> str:
    num = _num(status)
    if num is None:
        return _text(status)
    code = int(num)
    return _STATUS_LABELS.get(code, f"状态{code}")


def _has_quality_result(raw: dict) -> bool:
    """有质检完成时间，才进入下一步；扣重明细是否留存要再看详情。"""
    return bool(_text(raw.get("checkCompleteTime")))


def _positive(value: object) -> bool:
    number = _num(value)
    return number is not None and number > 0


def _empty_deduct_text(value: object) -> bool:
    text = _text(value)
    return not text or bool(_EMPTY_DEDUCT_TEXT.fullmatch(text))


def judgement_retained(deduction: dict) -> bool:
    """扣重为 0 且没有扣杂/厚度明细时，不算留下了检判记录。

    列表里的 deductWeight 不可用：有明细的车次列表扣重也经常是 0。
    这里只看详情里的总计和明细。拉详情失败时保留该车，避免把网络错误写成未检判。
    """
    if deduction.get("detail_error"):
        return True
    if _positive(deduction.get("total_ton")):
        return True
    if _positive(deduction.get("impurity_vis_ton")) or _positive(deduction.get("thick_vis_ton")):
        return True
    if _text(deduction.get("other_vis")):
        return True
    for key in ("impurity_vis_desc", "thick_vis_desc", "total_desc"):
        if not _empty_deduct_text(deduction.get(key)):
            return True
    return False


def load_manual_rows(path: str | Path) -> list[ManualRow]:
    """读取人工检判 xlsx。表头必须能认出「日期」和「车号」。"""
    file_path = Path(path)
    if not file_path.is_file():
        raise FileNotFoundError(f"找不到人工检判表: {file_path}")
    workbook = load_workbook(file_path, data_only=True)
    sheet = workbook.active
    header_row = next(sheet.iter_rows(min_row=1, max_row=1, values_only=True), None)
    if not header_row:
        raise ValueError(f"人工检判表是空的: {file_path}")
    columns: dict[str, int] = {}
    for index, header in enumerate(header_row):
        field_name = _HEADER_FIELDS.get(_text(header))
        if field_name and field_name not in columns:
            columns[field_name] = index
    missing = [name for name in ("date", "plate") if name not in columns]
    if missing:
        raise ValueError("人工检判表缺少「日期」或「车号」列，无法按车次对齐")

    rows: list[ManualRow] = []
    for excel_row, values in enumerate(sheet.iter_rows(min_row=2, values_only=True), start=2):
        if values is None or all(cell is None or _text(cell) == "" for cell in values):
            continue

        def cell(field_name: str) -> object:
            index = columns.get(field_name)
            if index is None or index >= len(values):
                return None
            return values[index]

        try:
            day = parse_manual_date(cell("date"))
        except ValueError as exc:
            raise ValueError(f"第 {excel_row} 行{exc}") from exc
        plate = _text(cell("plate"))
        rows.append(
            ManualRow(
                row_no=excel_row,
                date=day,
                plate=plate,
                plate_key=normalize_plate(plate),
                cargo=_text(cell("cargo")),
                supplier=_text(cell("supplier")),
                weight_ton=_num(cell("weight_ton")),
                user_unit=_text(cell("user_unit")),
                size_issue=_text(cell("size_issue")),
                size_ton=_num(cell("size_ton")),
                thick_issue=_text(cell("thick_issue")),
                thick_ton=_num(cell("thick_ton")),
                impurity_issue=_text(cell("impurity_issue")),
                impurity_ton=_num(cell("impurity_ton")),
                moisture_issue=_text(cell("moisture_issue")),
                moisture_ton=_num(cell("moisture_ton")),
                downgrade_issue=_text(cell("downgrade_issue")),
                downgrade_ton=_num(cell("downgrade_ton")),
                other_issue=_text(cell("other_issue")),
                other_ton=_num(cell("other_ton")),
                result=_text(cell("result")),
                deduct_ton=_num(cell("deduct_ton")),
            )
        )
    if not rows:
        raise ValueError(f"人工检判表没有数据行: {file_path}")
    return rows


def trip_from_raw(day: str, raw: dict) -> VisualTrip:
    plate = _text(raw.get("carNumber"))
    return VisualTrip(
        date=day,
        plate=plate,
        plate_key=normalize_plate(plate),
        flow_code=_text(raw.get("flowCode")),
        checked=_has_quality_result(raw),
        net_ton=kg_to_ton(raw.get("netWeight")),
        gross_ton=kg_to_ton(raw.get("grossWeight")),
        create_time=_text(raw.get("createTime")),
        raw=raw,
    )


def _weight_delta(manual: ManualRow, trip: VisualTrip) -> float:
    if manual.weight_ton is None or trip.net_ton is None:
        return 1e9
    return abs(manual.weight_ton - trip.net_ton)


def _matched_reason(visual_count: int, delta: float | None) -> str:
    if visual_count <= 1:
        text = "同日同车牌，唯一质检完成车次"
    else:
        text = f"同日同车牌有 {visual_count} 个质检完成车次，按净重就近配对"
    if delta is None:
        return text
    text = f"{text}，净重差 {delta:.3f} 吨"
    if abs(delta) >= 3:
        text += "，重量差较大请复核"
    return text


def match_day(
    manuals: list[ManualRow],
    visuals: list[VisualTrip],
) -> tuple[list[MatchResult], list[VisualTrip]]:
    """把同一天的人工行配到视觉车次。返回人工行结果，以及没配上的已完成视觉车次。"""
    ordered = sorted(manuals, key=lambda row: row.row_no)
    usable = [trip for trip in visuals if trip.checked and trip.retained and trip.plate_key]
    shell_plates = {trip.plate_key for trip in visuals if trip.checked and (not trip.retained) and trip.plate_key}
    if not ordered:
        return [], usable

    unchecked_plates = {trip.plate_key for trip in visuals if (not trip.checked) and trip.plate_key}
    if not usable:
        results: list[MatchResult] = []
        for row in ordered:
            if not row.plate_key:
                results.append(MatchResult(manual=row, status=STATUS_MISSING, reason_code=REASON_EMPTY_PLATE, reason="车号为空"))
            elif row.plate_key in shell_plates:
                results.append(MatchResult(
                    manual=row,
                    status=STATUS_MISSING,
                    reason_code=REASON_NO_DETAIL,
                    reason="视觉有车次，但扣重为0且没有扣杂/厚度检判明细",
                ))
            else:
                results.append(MatchResult(
                    manual=row,
                    status=STATUS_MISSING,
                    reason_code=REASON_NO_VISUAL_DAY,
                    reason="该日视觉系统没有质检完成记录",
                ))
        return results, []

    by_plate: dict[str, list[VisualTrip]] = defaultdict(list)
    for trip in usable:
        by_plate[trip.plate_key].append(trip)

    results: list[MatchResult] = []
    used: set[str] = set()
    for plate_key, rows in _groups(ordered):
        if not plate_key:
            for row in rows:
                results.append(
                    MatchResult(
                        manual=row,
                        status=STATUS_MISSING,
                        reason_code=REASON_EMPTY_PLATE,
                        reason="车号为空",
                    )
                )
            continue
        pool = by_plate.get(plate_key, [])
        if not pool:
            if plate_key in shell_plates:
                code = REASON_NO_DETAIL
                reason = "视觉有车次，但扣重为0且没有扣杂/厚度检判明细"
            elif plate_key in unchecked_plates:
                code = REASON_INCOMPLETE
                reason = "该车牌当天有过磅记录，但没有质检完成实绩"
            else:
                code = REASON_NO_PLATE
                reason = "视觉质检实绩中没有该车牌"
            for row in rows:
                results.append(MatchResult(manual=row, status=STATUS_MISSING, reason_code=code, reason=reason))
            continue

        pairs = [(row, trip) for row in rows for trip in pool]
        pairs.sort(key=lambda item: (_weight_delta(item[0], item[1]), item[0].row_no, item[1].create_time))
        assigned: dict[int, VisualTrip] = {}
        taken_flows: set[str] = set()
        for row, trip in pairs:
            if row.row_no in assigned or trip.flow_code in taken_flows:
                continue
            assigned[row.row_no] = trip
            taken_flows.add(trip.flow_code)
            used.add(trip.flow_code)
        for row in rows:
            trip = assigned.get(row.row_no)
            if trip is None:
                results.append(
                    MatchResult(
                        manual=row,
                        status=STATUS_MISSING,
                        reason_code=REASON_EXTRA_MANUAL,
                        reason="同车牌当天的质检车次已配给重量更接近的人工记录",
                    )
                )
                continue
            delta = None if row.weight_ton is None or trip.net_ton is None else _weight_delta(row, trip)
            results.append(
                MatchResult(
                    manual=row,
                    status=STATUS_ALIGNED,
                    reason_code=REASON_MATCHED,
                    reason=_matched_reason(len(pool), delta),
                    trip=trip,
                    deduction=dict(trip.deduction) if trip.deduction else {},
                )
            )
    leftovers = [trip for trip in usable if trip.flow_code not in used]
    results.sort(key=lambda item: item.manual.row_no)
    return results, leftovers


def _groups(rows: list[ManualRow]) -> list[tuple[str, list[ManualRow]]]:
    grouped: dict[str, list[ManualRow]] = defaultdict(list)
    order: list[str] = []
    for row in rows:
        if row.plate_key not in grouped:
            order.append(row.plate_key)
        grouped[row.plate_key].append(row)
    return [(key, grouped[key]) for key in order]


def _merge_items(items: list[dict]) -> tuple[float | None, float | None, str]:
    if not items:
        return None, None, ""
    weights = [_num(item.get("deductionWeight")) for item in items]
    rates = [_num(item.get("deductionRate")) for item in items]
    weight = sum(item for item in weights if item is not None) if any(item is not None for item in weights) else None
    rate = sum(item for item in rates if item is not None) if any(item is not None for item in rates) else None
    descriptions = [_text(item.get("description")) for item in items]
    return weight, rate, "；".join(item for item in descriptions if item)


def _split_items(items: list[dict]) -> dict[str, list[dict]]:
    buckets: dict[str, list[dict]] = {"IMPURITY": [], "THICKNESS": [], "OTHER": []}
    for item in items:
        code = _text(item.get("deductionType")).upper()
        name = _text(item.get("deductionTypeName"))
        if code == "IMPURITY" or name == "扣杂":
            buckets["IMPURITY"].append(item)
        elif code == "THICKNESS" or "厚度" in name:
            buckets["THICKNESS"].append(item)
        else:
            buckets["OTHER"].append(item)
    return buckets


def _detail_payload(detail: dict) -> tuple[dict, str]:
    """优先用整车总计。总计没有扣重体时，才退回唯一工位的明细。"""
    total = ((detail.get("totalCheckResult") or {}).get("deductionDetail") or {})
    if total.get("deductionItems") or total.get("deductionDescription") or total.get("totalDeductionWeight") not in (None, ""):
        return total, "整车总计"
    station_details = [
        (station.get("deductionDetail") or {})
        for station in (detail.get("stationResults") or [])
        if station.get("deductionDetail")
    ]
    usable = [
        item for item in station_details
        if item.get("deductionItems") or item.get("deductionDescription") or item.get("totalDeductionWeight") not in (None, "")
    ]
    if len(usable) == 1:
        return usable[0], "单工位"
    return total, "整车总计"


def parse_deduction(detail: dict) -> dict:
    """把扣重扣杂结果拆成总计、扣杂、厚度扣重。重量单位沿用接口里的吨。"""
    payload, source = _detail_payload(detail or {})
    buckets = _split_items(list(payload.get("deductionItems") or []))
    impurity_ton, impurity_rate, impurity_desc = _merge_items(buckets["IMPURITY"])
    thick_ton, thick_rate, thick_desc = _merge_items(buckets["THICKNESS"])
    other_parts = []
    for item in buckets["OTHER"]:
        name = _text(item.get("deductionTypeName")) or _text(item.get("deductionType")) or "其他"
        weight, rate, desc = _merge_items([item])
        bits = [name]
        if weight is not None:
            bits.append(f"{weight}吨")
        if rate is not None:
            bits.append(f"{rate}%")
        if desc:
            bits.append(desc)
        other_parts.append(" ".join(bits))
    return {
        "total_ton": _num(payload.get("totalDeductionWeight")),
        "total_rate": _num(payload.get("deductionRate")),
        "total_desc": _text(payload.get("deductionDescription")),
        "standard": _text(payload.get("calculationStandard")),
        "declared_name": _text(payload.get("declaredSteelTypeName")),
        "impurity_vis_ton": impurity_ton,
        "impurity_vis_rate": impurity_rate,
        "impurity_vis_desc": impurity_desc,
        "thick_vis_ton": thick_ton,
        "thick_vis_rate": thick_rate,
        "thick_vis_desc": thick_desc,
        "other_vis": "；".join(other_parts),
        "detail_source": source,
        "detail_error": "",
    }


def _output_path(rows: list[ManualRow], output: str | None) -> Path:
    start = min(row.date for row in rows)
    end = max(row.date for row in rows)
    name = f"永锋废钢扣杂准确率_{start}_{end}.xlsx"
    if not output:
        path = Path("downloads") / "yongfeng" / name
    else:
        path = Path(output)
        if path.suffix.lower() != ".xlsx":
            path = path / name
    path.parent.mkdir(parents=True, exist_ok=True)
    return path


def _row_dict(item: MatchResult) -> dict:
    manual = item.manual
    trip = item.trip
    deduction = item.deduction or {}
    raw = trip.raw if trip else {}
    aligned = item.status == STATUS_ALIGNED
    total_ton = deduction.get("total_ton") if aligned else None
    impurity_vis = deduction.get("impurity_vis_ton") if aligned else None
    thick_vis = deduction.get("thick_vis_ton") if aligned else None
    deduct_diff = _sub(total_ton, manual.deduct_ton) if aligned else None
    return {
        "row_no": manual.row_no,
        "date": manual.date,
        "plate": manual.plate,
        "cargo": manual.cargo,
        "supplier": manual.supplier,
        "weight_ton": manual.weight_ton,
        "user_unit": manual.user_unit,
        "size_issue": manual.size_issue,
        "size_ton": manual.size_ton,
        "thick_issue": manual.thick_issue,
        "thick_ton": manual.thick_ton,
        "impurity_issue": manual.impurity_issue,
        "impurity_ton": manual.impurity_ton,
        "moisture_issue": manual.moisture_issue,
        "moisture_ton": manual.moisture_ton,
        "downgrade_issue": manual.downgrade_issue,
        "downgrade_ton": manual.downgrade_ton,
        "other_issue": manual.other_issue,
        "other_ton": manual.other_ton,
        "result": manual.result,
        "deduct_ton": manual.deduct_ton,
        "status": item.status,
        "reason": item.reason if not deduction.get("detail_error") else f"{item.reason}；扣重明细获取失败：{deduction.get('detail_error')}",
        "bill_no": _text(raw.get("weightBillNo")) if aligned else "",
        "flow_code": trip.flow_code if aligned and trip else "",
        "station": raw.get("stationNumber") if aligned else "",
        "visual_status": _status_label(raw.get("status")) if aligned else "",
        "create_time": trip.create_time if aligned and trip else "",
        "complete_time": _text(raw.get("checkCompleteTime")) if aligned else "",
        "net_ton": trip.net_ton if aligned and trip else None,
        "gross_ton": trip.gross_ton if aligned and trip else None,
        "material_name": _text(raw.get("materialName")) if aligned else "",
        "declared_name": deduction.get("declared_name", "") if aligned else "",
        "total_ton": total_ton,
        "total_rate": deduction.get("total_rate") if aligned else None,
        "total_desc": deduction.get("total_desc", "") if aligned else "",
        "standard": deduction.get("standard", "") if aligned else "",
        "impurity_vis_ton": impurity_vis,
        "impurity_vis_rate": deduction.get("impurity_vis_rate") if aligned else None,
        "impurity_vis_desc": deduction.get("impurity_vis_desc", "") if aligned else "",
        "thick_vis_ton": thick_vis,
        "thick_vis_rate": deduction.get("thick_vis_rate") if aligned else None,
        "thick_vis_desc": deduction.get("thick_vis_desc", "") if aligned else "",
        "other_vis": deduction.get("other_vis", "") if aligned else "",
        "deduct_diff": deduct_diff,
        "deduct_diff_kg": None if deduct_diff is None else round(deduct_diff * 1000, 1),
        "verdict": judge_deduct_diff(deduct_diff),
        "impurity_diff": _sub(impurity_vis, manual.impurity_ton) if aligned else None,
        "thick_diff": _sub(thick_vis, manual.thick_ton) if aligned else None,
    }


def _apply_header(sheet, columns: list[tuple[str, str, str, int, str]]) -> None:
    thin = Border(
        left=Side(style="thin", color="D0D7DE"),
        right=Side(style="thin", color="D0D7DE"),
        top=Side(style="thin", color="D0D7DE"),
        bottom=Side(style="thin", color="D0D7DE"),
    )
    start = 1
    while start <= len(columns):
        group = columns[start - 1][0]
        end = start
        while end < len(columns) and columns[end][0] == group:
            end += 1
        color = _GROUP_FILLS.get(group, "F5F5F5")
        sheet.merge_cells(start_row=1, start_column=start, end_row=1, end_column=end)
        for col in range(start, end + 1):
            cell = sheet.cell(1, col, group if col == start else None)
            cell.fill = PatternFill("solid", fgColor=color)
            cell.font = Font(bold=True, color="1F2328")
            cell.alignment = Alignment(horizontal="center", vertical="center")
            cell.border = thin
        start = end + 1
    for index, (_group, title, _key, width, _kind) in enumerate(columns, start=1):
        cell = sheet.cell(2, index, title)
        cell.font = Font(bold=True)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = thin
        cell.fill = PatternFill("solid", fgColor="F6F8FA")
        sheet.column_dimensions[get_column_letter(index)].width = width
    sheet.row_dimensions[1].height = 22
    sheet.row_dimensions[2].height = 32
    sheet.freeze_panes = "A3"
    sheet.auto_filter.ref = f"A2:{get_column_letter(len(columns))}2"


def _text_row_height(text: str, width: int) -> float:
    """按列宽估算换行后的行高。厚度扣重说明往往是一整段公式。"""
    if not text:
        return 18
    units = sum(2 if ord(ch) > 127 else 1 for ch in text)
    per_line = max(8, int(width))
    lines = max(1, (units + per_line - 1) // per_line)
    return float(min(140, max(18, lines * 16)))


def _write_body(sheet, columns: list[tuple[str, str, str, int, str]], records: list[dict]) -> None:
    thin = Border(
        left=Side(style="thin", color="D0D7DE"),
        right=Side(style="thin", color="D0D7DE"),
        top=Side(style="thin", color="D0D7DE"),
        bottom=Side(style="thin", color="D0D7DE"),
    )
    aligned_fill = PatternFill("solid", fgColor="C8E6C9")
    missing_fill = PatternFill("solid", fgColor="E0E0E0")
    pass_fill = PatternFill("solid", fgColor="A5D6A7")
    fail_fill = PatternFill("solid", fgColor="EF9A9A")
    thick_width = next((width for _group, _title, key, width, _kind in columns if key == "thick_vis_desc"), 48)
    for row_index, record in enumerate(records, start=3):
        verdict = record.get("verdict") or ""
        for col_index, (_group, _title, key, _width, kind) in enumerate(columns, start=1):
            value = record.get(key, "")
            cell = sheet.cell(row_index, col_index, None if value == "" else value)
            cell.border = thin
            wrap = key in _WRAP_KEYS
            cell.alignment = Alignment(vertical="center", wrap_text=wrap, horizontal="center" if key in {"status", "verdict"} else "left")
            if kind == "ton" and isinstance(value, (int, float)):
                cell.number_format = "0.0000"
            elif kind == "kg" and isinstance(value, (int, float)):
                cell.number_format = "0.0"
            elif kind == "rate" and isinstance(value, (int, float)):
                cell.number_format = "0.00"
            if key == "status":
                cell.fill = aligned_fill if value == STATUS_ALIGNED else missing_fill
                cell.font = Font(bold=True)
            elif key in _DIFF_KEYS and verdict == "合格":
                cell.fill = pass_fill
                cell.font = Font(bold=True, color="1B5E20")
            elif key in _DIFF_KEYS and verdict == "不合格":
                cell.fill = fail_fill
                cell.font = Font(bold=True, color="B71C1C")
        sheet.row_dimensions[row_index].height = _text_row_height(str(record.get("thick_vis_desc") or ""), thick_width)
    last_row = 2 + len(records)
    sheet.auto_filter.ref = f"A2:{get_column_letter(len(columns))}{last_row}"


def write_alignment_workbook(
    aligned: list[MatchResult],
    visual_only: list[tuple[str, VisualTrip]],
    summary: dict,
    path: Path,
) -> None:
    """写出对齐明细、视觉有人工无、汇总三张表。"""
    workbook = Workbook()
    detail = workbook.active
    detail.title = "对齐明细"
    _apply_header(detail, _COLUMNS)
    _write_body(detail, _COLUMNS, [_row_dict(item) for item in aligned])

    extra = workbook.create_sheet("视觉有人工无")
    extra_columns = [
        ("视觉有人工无", "日期", "date", 12, "text"),
        ("视觉有人工无", "车牌", "plate", 14, "text"),
        ("视觉有人工无", "净重（吨）", "net_ton", 12, "ton"),
        ("视觉有人工无", "说明", "reason", 42, "text"),
    ]
    _apply_header(extra, extra_columns)
    extra_rows = []
    for reason, trip in visual_only:
        extra_rows.append({
            "date": trip.date,
            "plate": trip.plate,
            "net_ton": trip.net_ton,
            "reason": reason,
        })
    _write_body(extra, extra_columns, extra_rows)

    cover = workbook.create_sheet("汇总", 0)
    cover["A1"] = "永锋废钢视觉质检实绩 × 人工检判对齐"
    cover["A1"].font = Font(bold=True, size=14)
    cover["A3"] = (
        "本表一个统计周期对应人工表的起止日期。扣杂准确率只统计已对齐、双方扣重都有数的车："
        "视觉总计重量减去人工扣除重量，绝对值在 150kg 以内为合格，否则不合格。"
        "未检判不进入准确率。绿色是合格，红色是不合格。"
        "扣重为 0 且没有扣杂/厚度明细的视觉车次已排除。"
    )
    cover["A3"].alignment = Alignment(wrap_text=True, vertical="center")
    cover.row_dimensions[3].height = 62
    cover.merge_cells("A3:B3")
    rate_text = summary.get("deduct_rate_text") or "—"
    labels = [
        ("统计周期", f"{summary.get('date_start') or ''} 至 {summary.get('date_end') or ''}"),
        ("扣杂准确率", rate_text),
        ("合格车数（|差值|≤150kg）", summary.get("deduct_ok")),
        ("不合格车数", summary.get("deduct_bad")),
        ("可评估车数", summary.get("deduct_evaluable")),
        ("人工检判行数", summary.get("manual_rows")),
        ("已对齐", summary.get("aligned")),
        ("未检判", summary.get("unjudged")),
        ("未检判·该日无质检完成记录", summary.get("no_visual_day")),
        ("未检判·车号为空", summary.get("empty_plate")),
        ("未检判·视觉无此车牌", summary.get("no_plate")),
        ("未检判·有过磅但质检未完成", summary.get("incomplete")),
        ("未检判·扣重为0且无检判明细", summary.get("no_detail")),
        ("未检判·同车牌车次已配给其他行", summary.get("extra_manual")),
        ("视觉车次中扣重为0已排除", summary.get("shell_trips")),
        ("视觉质检完成车次", summary.get("visual_checked")),
        ("视觉有、人工表没有的车次", summary.get("visual_only")),
        ("人工日期起", summary.get("date_start")),
        ("人工日期止", summary.get("date_end")),
        ("人工表", summary.get("source")),
        ("生成时间", summary.get("generated_at")),
    ]
    for index, (label, value) in enumerate(labels, start=5):
        cover.cell(index, 1, label).font = Font(bold=True)
        cell = cover.cell(index, 2, value)
        if label == "扣杂准确率":
            cell.font = Font(bold=True, size=16, color="1B5E20")
    cover.column_dimensions["A"].width = 36
    cover.column_dimensions["B"].width = 80
    workbook.save(path)


def _summarize(
    aligned: list[MatchResult],
    visual_checked: int,
    visual_only: int,
    shell_trips: int,
    source: str,
) -> dict:
    counts = defaultdict(int)
    for item in aligned:
        counts[item.reason_code] += 1
    dates = [item.manual.date for item in aligned]
    verdicts = [(_row_dict(item).get("verdict") or "") for item in aligned]
    deduct_ok = sum(1 for verdict in verdicts if verdict == "合格")
    deduct_bad = sum(1 for verdict in verdicts if verdict == "不合格")
    deduct_evaluable = deduct_ok + deduct_bad
    deduct_rate = (deduct_ok / deduct_evaluable) if deduct_evaluable else None
    return {
        "manual_rows": len(aligned),
        "aligned": sum(1 for item in aligned if item.status == STATUS_ALIGNED),
        "unjudged": sum(1 for item in aligned if item.status == STATUS_MISSING),
        "no_visual_day": counts[REASON_NO_VISUAL_DAY],
        "empty_plate": counts[REASON_EMPTY_PLATE],
        "no_plate": counts[REASON_NO_PLATE],
        "incomplete": counts[REASON_INCOMPLETE],
        "no_detail": counts[REASON_NO_DETAIL],
        "extra_manual": counts[REASON_EXTRA_MANUAL],
        "shell_trips": shell_trips,
        "visual_checked": visual_checked,
        "visual_only": visual_only,
        "deduct_ok": deduct_ok,
        "deduct_bad": deduct_bad,
        "deduct_evaluable": deduct_evaluable,
        "deduct_rate": deduct_rate,
        "deduct_rate_text": "—" if deduct_rate is None else f"{deduct_rate * 100:.2f}%",
        "date_start": min(dates) if dates else "",
        "date_end": max(dates) if dates else "",
        "source": source,
        "generated_at": datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
    }


def format_align_summary(result: dict) -> str:
    """给对话用的短摘要。数字全部来自本次对齐结果。"""
    summary = result.get("summary") or {}
    return (
        f"已生成永锋废钢扣杂准确率表（{summary.get('date_start', '')} 至 {summary.get('date_end', '')}）\n"
        f"扣杂准确率 {summary.get('deduct_rate_text', '—')}："
        f"合格 {summary.get('deduct_ok', 0)} / 可评估 {summary.get('deduct_evaluable', 0)}，"
        f"不合格 {summary.get('deduct_bad', 0)}（|视觉总计−人工扣除|≤150kg 为合格）\n"
        f"人工记录 {summary.get('manual_rows', 0)} 行，"
        f"已对齐 {summary.get('aligned', 0)} 行，"
        f"未检判 {summary.get('unjudged', 0)} 行\n"
        f"其中：该日无质检完成记录 {summary.get('no_visual_day', 0)}，"
        f"车号为空 {summary.get('empty_plate', 0)}，"
        f"视觉无此车牌 {summary.get('no_plate', 0)}，"
        f"有过磅但质检未完成 {summary.get('incomplete', 0)}，"
        f"扣重为0且无检判明细 {summary.get('no_detail', 0)}，"
        f"同车牌车次已配给其他行 {summary.get('extra_manual', 0)}\n"
        f"视觉车次里因扣重为0、没有明细而排除 {summary.get('shell_trips', 0)} 趟\n"
        f"这些日期里视觉质检完成 {summary.get('visual_checked', 0)} 车，"
        f"其中人工表没有的 {summary.get('visual_only', 0)} 车（见「视觉有人工无」）\n"
        f"文件：{result.get('xlsx_path', '')}"
    )


def _query_day(client: YongfengScrapClient, day: str):
    try:
        return client.query_list_by_date(day)
    except Exception:
        logger.warning("永锋质检实绩 %s 查询失败，重试一次", day, exc_info=True)
        return client.query_list_by_date(day)


def _annotate_retained(client: YongfengScrapClient, trips: list[VisualTrip]) -> int:
    """扣重为 0 且没有明细的已完成车次，标记为不参与对齐。"""
    dropped = 0
    for trip in trips:
        if not trip.checked or not trip.flow_code:
            continue
        try:
            detail = client.get_detail_by_flow(trip.flow_code, trip.raw.get("stationNumber"))
            trip.deduction = parse_deduction(detail)
        except Exception as exc:
            logger.warning("扣重明细失败 %s %s: %s", trip.plate, trip.flow_code, exc)
            trip.deduction = {"detail_error": str(exc), "detail_source": ""}
            trip.retained = True
            continue
        trip.retained = judgement_retained(trip.deduction)
        if not trip.retained:
            dropped += 1
    return dropped


def _fill_deduction(client: YongfengScrapClient, item: MatchResult) -> None:
    trip = item.trip
    if item.status != STATUS_ALIGNED or trip is None or not trip.flow_code:
        return
    if trip.deduction:
        item.deduction = dict(trip.deduction)
        return
    try:
        detail = client.get_detail_by_flow(trip.flow_code, trip.raw.get("stationNumber"))
        item.deduction = parse_deduction(detail)
    except Exception as exc:
        logger.warning("扣重明细失败 %s %s: %s", trip.plate, trip.flow_code, exc)
        item.deduction = {"detail_error": str(exc), "detail_source": ""}


def iter_align_manual(
    manual_file: str,
    output: str | None = None,
) -> Iterator[dict]:
    """按人工表里的日期逐天拉视觉质检实绩，并持续吐出进度。"""
    source = str(Path(manual_file))
    manuals = load_manual_rows(source)
    by_date: dict[str, list[ManualRow]] = defaultdict(list)
    for row in manuals:
        by_date[row.date].append(row)
    dates = sorted(by_date)
    yield {
        "type": "status",
        "message": f"人工表 {len(manuals)} 行，{dates[0]} 至 {dates[-1]}，共 {len(dates)} 天。开始登录永锋质检实绩。",
    }
    aligned: list[MatchResult] = []
    visual_only: list[tuple[str, VisualTrip]] = []
    visual_checked = 0
    shell_trips = 0
    with YongfengScrapClient() as client:
        client.login()
        for index, day in enumerate(dates, start=1):
            records = _query_day(client, day)
            trips = [trip_from_raw(day, record.raw) for record in records]
            day_shells = _annotate_retained(client, trips)
            shell_trips += day_shells
            day_kept = sum(1 for trip in trips if trip.checked and trip.retained)
            visual_checked += day_kept
            matched, leftovers = match_day(by_date[day], trips)
            for item in matched:
                _fill_deduction(client, item)
            aligned.extend(matched)
            manual_plates = {row.plate_key for row in by_date[day] if row.plate_key}
            for trip in leftovers:
                if trip.plate_key in manual_plates:
                    reason = "同车牌人工记录已配完，这趟视觉车次没有对应人工行"
                else:
                    reason = "人工表无此车牌"
                visual_only.append((reason, trip))
            hit = sum(1 for item in matched if item.status == STATUS_ALIGNED)
            yield {
                "type": "status",
                "message": (
                    f"[{index}/{len(dates)}] {day} 人工 {len(matched)} 行，"
                    f"有效检判 {day_kept} 车，扣重为0已排除 {day_shells} 车，对齐 {hit} 行"
                ),
            }
    aligned.sort(key=lambda item: item.manual.row_no)
    summary = _summarize(aligned, visual_checked, len(visual_only), shell_trips, source)
    path = _output_path(manuals, output)
    write_alignment_workbook(aligned, visual_only, summary, path)
    yield {
        "type": "done",
        "result": {
            "xlsx_path": str(path),
            "summary": summary,
        },
    }


def align_manual_file(
    manual_file: str,
    output: str | None = None,
    progress: Callable[[str], None] | None = None,
) -> dict:
    """跑完对齐并返回 xlsx 路径和汇总。不连网的测试不要调用它。"""
    result = None
    for event in iter_align_manual(manual_file, output=output):
        if event.get("type") == "status" and progress is not None:
            progress(str(event.get("message") or ""))
        if event.get("type") == "done":
            result = event.get("result")
    if not result:
        raise RuntimeError("对齐没有产出表格")
    return result


def main() -> None:
    logging.basicConfig(level=logging.INFO, format="%(asctime)s [%(name)s] %(levelname)s: %(message)s")
    parser = argparse.ArgumentParser(description="永锋废钢视觉质检实绩与人工检判表对齐")
    parser.add_argument("--manual", required=True, help="本地人工检判 xlsx 绝对路径")
    parser.add_argument("--output", default="", help="输出 xlsx 或目录；留空则写到 downloads/yongfeng")
    args = parser.parse_args()
    result = align_manual_file(args.manual, output=args.output or None, progress=print)
    print(format_align_summary(result))


if __name__ == "__main__":
    main()
