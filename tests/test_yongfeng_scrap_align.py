"""永锋废钢视觉/人工对齐：不连视觉系统。

运行：python tests/test_yongfeng_scrap_align.py
"""
from __future__ import annotations

import sys
from datetime import datetime
from pathlib import Path
from tempfile import TemporaryDirectory

from openpyxl import Workbook, load_workbook

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

from agent.tools import TOOLS
from agent.yongfeng.scrap_align import (
    ManualRow,
    MatchResult,
    VisualTrip,
    is_yongfeng_manual_align,
    judgement_retained,
    judge_deduct_diff,
    load_manual_rows,
    match_day,
    parse_deduction,
    parse_manual_date,
    resolve_manual_path,
    write_alignment_workbook,
    _summarize,
)


def _manual(row_no: int, day: str, plate: str, weight: float) -> ManualRow:
    return ManualRow(
        row_no=row_no,
        date=day,
        plate=plate,
        plate_key=plate.upper(),
        weight_ton=weight,
        impurity_ton=0.3,
        deduct_ton=0.3,
    )


def _trip(day: str, plate: str, flow: str, net: float, checked: bool) -> VisualTrip:
    return VisualTrip(
        date=day,
        plate=plate,
        plate_key=plate.upper(),
        flow_code=flow,
        checked=checked,
        net_ton=net,
        gross_ton=net + 10,
        create_time=f"{day} 08:00:00",
        raw={"weightBillNo": "B1", "checkCompleteTime": f"{day} 09:00:00" if checked else "", "status": 4},
    )


def test_excel_serial_and_text_dates():
    assert parse_manual_date(46235) == "2026-08-01"
    assert parse_manual_date(datetime(2026, 8, 2)) == "2026-08-02"
    assert parse_manual_date("2026/8/3") == "2026-08-03"


def test_match_nearest_weight_and_missing_reasons():
    manuals = [
        _manual(1, "2026-10-05", "苏GQ1389", 43.60),
        _manual(2, "2026-10-05", "苏GQ1389", 43.28),
        _manual(3, "2026-10-05", "鲁A00001", 10),
        _manual(4, "2026-10-05", "鲁C00001", 20),
    ]
    visuals = [
        _trip("2026-10-05", "苏GQ1389", "f1", 43.10, True),
        _trip("2026-10-05", "苏GQ1389", "f2", 43.55, True),
        _trip("2026-10-05", "鲁B99999", "f3", 20, True),
        _trip("2026-10-05", "鲁A00001", "f4", 10, False),
    ]
    matched, leftover = match_day(manuals, visuals)
    assert matched[0].trip is not None and matched[0].trip.flow_code == "f2"
    assert matched[1].trip is not None and matched[1].trip.flow_code == "f1"
    assert matched[2].status == "未检判" and matched[2].reason_code == "incomplete"
    assert matched[3].status == "未检判" and matched[3].reason_code == "no_plate"
    assert [trip.flow_code for trip in leftover] == ["f3"]


def test_zero_deduction_without_items_is_unjudged():
    shell = _trip("2026-09-06", "冀AZC799", "empty", 41.0, True)
    shell.retained = False
    matched, leftover = match_day([_manual(8, "2026-09-06", "冀AZC799", 41.0)], [shell])
    assert matched[0].status == "未检判"
    assert matched[0].reason_code == "no_detail"
    assert leftover == []

    real = _trip("2026-09-06", "冀AZC799", "real", 40.0, True)
    matched, _leftover = match_day([_manual(8, "2026-09-06", "冀AZC799", 41.0)], [shell, real])
    assert matched[0].status == "已对齐"
    assert matched[0].trip is not None and matched[0].trip.flow_code == "real"


def test_judgement_retained_rejects_empty_shell():
    shell = {
        "total_ton": 0.0,
        "total_desc": "0.00吨（）",
        "impurity_vis_ton": None,
        "thick_vis_ton": None,
        "impurity_vis_desc": "",
        "thick_vis_desc": "",
        "other_vis": "",
    }
    assert judgement_retained(shell) is False
    assert judgement_retained({"total_ton": 0.243, "impurity_vis_ton": 0.0543, "thick_vis_ton": 0.1887}) is True
    assert judgement_retained({"detail_error": "timeout", "total_ton": 0}) is True


def test_day_without_visual_is_unjudged():
    matched, leftover = match_day([_manual(2, "2026-08-01", "鲁NK1805", 39.72)], [])
    assert matched[0].status == "未检判"
    assert matched[0].reason_code == "no_visual_day"
    assert leftover == []


def test_parse_deduction_splits_impurity_and_thickness():
    detail = {
        "totalCheckResult": {
            "deductionDetail": {
                "totalDeductionWeight": 0.243,
                "deductionDescription": "0.24吨（扣杂0.05吨｜厚度扣重0.19吨）",
                "deductionRate": 0.52,
                "calculationStandard": "厚度低于10mm的超过部分折扣",
                "declaredSteelTypeName": "重废1",
                "deductionItems": [
                    {
                        "deductionType": "IMPURITY",
                        "deductionTypeName": "扣杂",
                        "deductionWeight": 0.0543,
                        "deductionRate": 0.12,
                        "description": "扣杂占比0.1%｜0.05吨",
                    },
                    {
                        "deductionType": "THICKNESS",
                        "deductionTypeName": "厚度扣重",
                        "deductionWeight": 0.189,
                        "deductionRate": 0.40,
                        "description": "厚度扣重0.19吨",
                    },
                ],
            }
        }
    }
    parsed = parse_deduction(detail)
    assert parsed["total_ton"] == 0.243
    assert parsed["impurity_vis_ton"] == 0.0543
    assert parsed["thick_vis_ton"] == 0.189
    assert parsed["detail_source"] == "整车总计"
    assert "扣杂" in parsed["impurity_vis_desc"]


def test_workbook_marks_unjudged_and_keeps_manual_row():
    item_day = match_day([_manual(5, "2026-08-01", "鲁NK1805", 39.72)], [])[0][0]
    summary = {
        "manual_rows": 1,
        "aligned": 0,
        "unjudged": 1,
        "no_visual_day": 1,
        "empty_plate": 0,
        "no_plate": 0,
        "incomplete": 0,
        "extra_manual": 0,
        "visual_checked": 0,
        "visual_only": 0,
        "date_start": "2026-08-01",
        "date_end": "2026-08-01",
        "source": "demo.xlsx",
        "generated_at": "2026-10-27 13:00:00",
    }
    with TemporaryDirectory() as tmp:
        path = Path(tmp) / "out.xlsx"
        write_alignment_workbook([item_day], [], summary, path)
        book = load_workbook(path)
        sheet = book["对齐明细"]
        headers = [sheet.cell(2, col).value for col in range(1, sheet.max_column + 1)]
        assert sheet["A1"].value == "人工检判"
        assert sheet.cell(3, headers.index("对齐状态") + 1).value == "未检判"
        assert sheet.cell(3, headers.index("日期") + 1).value == "2026-08-01"
        assert "扣杂重量（吨）" in headers
        assert "扣重判定" in headers
        for banned in ("磅单号", "车次号", "工位", "状态", "创建时间", "质检完成时间", "明细来源"):
            assert banned not in headers
        thick_col = headers.index("厚度扣重说明") + 1
        assert sheet.cell(3, thick_col).alignment.wrap_text is True


def test_150kg_splits_pass_and_fail_for_one_period():
    assert judge_deduct_diff(0.15) == "合格"
    assert judge_deduct_diff(-0.15) == "合格"
    assert judge_deduct_diff(0.151) == "不合格"
    assert judge_deduct_diff(None) == ""
    passed = _manual(5, "2026-08-07", "晋D03800D", 32.72)
    passed.deduct_ton = 0.50
    failed = _manual(6, "2026-08-07", "晋D00001", 30)
    failed.deduct_ton = 0.50
    rows = [
        MatchResult(manual=passed, status="已对齐", reason_code="matched", reason="", deduction={"total_ton": 0.60}),
        MatchResult(manual=failed, status="已对齐", reason_code="matched", reason="", deduction={"total_ton": 0.80}),
    ]
    summary = _summarize(rows, 2, 0, 0, "demo.xlsx")
    assert summary["deduct_ok"] == 1
    assert summary["deduct_bad"] == 1
    assert summary["deduct_evaluable"] == 2
    assert summary["deduct_rate_text"] == "50.00%"
    assert summary["date_start"] == "2026-08-07"


def test_load_real_manual_headers():
    path = Path(r"D:\BriefMe_Agent\废钢人工判定.xlsx")
    if not path.is_file():
        return
    rows = load_manual_rows(path)
    assert len(rows) == 422
    assert rows[0].date == "2026-08-01"
    assert rows[0].plate == "鲁NK1805"
    assert rows[0].deduct_ton == 0.3
    assert rows[-1].date == "2026-10-21"


def test_route_and_path_and_tool():
    assert is_yongfeng_manual_align("对齐【永锋】废钢视觉质检实绩和人工检判结果")
    assert not is_yongfeng_manual_align("下载 2026-09-01 的【永锋】检判原图")
    assert not is_yongfeng_manual_align("生成 2026-04-01 到 2026-04-07 的烧结矿颗粒度准确率报表")
    assert not is_yongfeng_manual_align("对齐盛隆人工检判结果")
    message = r"对齐【永锋】废钢 D:\BriefMe_Agent\废钢人工判定.xlsx"
    assert resolve_manual_path(message, r"D:\other.xlsx").endswith("废钢人工判定.xlsx")
    assert resolve_manual_path("对齐【永锋】废钢视觉质检实绩和人工检判结果", r"D:\picked.xlsx") == r"D:\picked.xlsx"
    names = [item["function"]["name"] for item in TOOLS]
    assert "yongfeng_align_scrap_manual" in names


def _write_sample_book(path: Path) -> None:
    book = Workbook()
    sheet = book.active
    sheet.append(["日期", "车号", "货物名称", "重量（吨）", "非钢杂质", "非钢杂质扣吨", "扣除重量（吨）"])
    sheet.append([46235, "鲁NK1805", "统料", 39.72, "杂质", 0.3, 0.3])
    book.save(path)


def test_load_sample_workbook():
    with TemporaryDirectory() as tmp:
        path = Path(tmp) / "manual.xlsx"
        _write_sample_book(path)
        rows = load_manual_rows(path)
        assert rows[0].date == "2026-08-01"
        assert rows[0].impurity_ton == 0.3


if __name__ == "__main__":
    test_excel_serial_and_text_dates()
    test_match_nearest_weight_and_missing_reasons()
    test_zero_deduction_without_items_is_unjudged()
    test_judgement_retained_rejects_empty_shell()
    test_day_without_visual_is_unjudged()
    test_parse_deduction_splits_impurity_and_thickness()
    test_workbook_marks_unjudged_and_keeps_manual_row()
    test_150kg_splits_pass_and_fail_for_one_period()
    test_load_sample_workbook()
    test_load_real_manual_headers()
    test_route_and_path_and_tool()
    print("ok")
